#!/usr/bin/env python3
"""Build a reviewable MTPC knowledge bundle from an office-document folder."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path


SUPPORTED = {".docx", ".pptx", ".xlsx", ".pdf", ".txt", ".md", ".csv"}
PRIVATE_FILE_PATTERNS = ("thẻ học sinh", "the hoc sinh", "untitled_m2")
ACTIVE_DOCUMENTS = {
    "doc1.docx",
    "phiếu đăng kí dự tuyển.docx",
    "to_roi_1.docx",
    "ts1.docx",
    "ts2.docx",
    "thư ngo_2024.docx",
}


def clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    blocks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return clean("\n".join(blocks))


def read_pptx(path: Path) -> str:
    from pptx import Presentation

    deck = Presentation(path)
    blocks = []
    for index, slide in enumerate(deck.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            blocks.append("[Trang chiếu %d]\n%s" % (index, "\n".join(slide_text)))
    return clean("\n".join(blocks))


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for sheet in workbook.worksheets:
        blocks.append("[Bảng: %s]" % sheet.title)
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                blocks.append(" | ".join(values))
    workbook.close()
    return clean("\n".join(blocks))


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return clean("\n".join((page.extract_text() or "") for page in reader.pages))


def read_text(path: Path) -> str:
    return clean(path.read_text(encoding="utf-8-sig", errors="replace"))


def extract(path: Path) -> str:
    readers = {
        ".docx": read_docx,
        ".pptx": read_pptx,
        ".xlsx": read_xlsx,
        ".pdf": read_pdf,
        ".txt": read_text,
        ".md": read_text,
        ".csv": read_text,
    }
    return readers[path.suffix.lower()](path)


def infer_year(path: Path, text: str) -> int:
    candidates = [int(value) for value in re.findall(r"\b20(?:1[5-9]|2[0-9])\b", path.name + "\n" + text[:4000])]
    if candidates:
        return max(candidates)
    return datetime.fromtimestamp(path.stat().st_mtime).year


def infer_type(name: str, text: str) -> str:
    value = (name + " " + text[:1200]).lower()
    if "phiếu đăng" in value or "phieu dang" in value:
        return "Biểu mẫu tuyển sinh"
    if "thư ngỏ" in value or "thu ngo" in value:
        return "Thư ngỏ"
    if "tuyển sinh" in value or re.search(r"\bts[_ -]?20", value):
        return "Tuyển sinh"
    if "thẻ học sinh" in value:
        return "Biểu mẫu học sinh"
    return "Thông tin trường"


def chunks_for(source: dict, size: int = 1100, overlap: int = 180) -> list[dict]:
    text = source["text"]
    result = []
    offset = 0
    number = 0
    while offset < len(text):
        end = min(len(text), offset + size)
        if end < len(text):
            break_at = max(text.rfind("\n", offset, end), text.rfind(". ", offset, end))
            if break_at > offset + size // 2:
                end = break_at + 1
        part = text[offset:end].strip()
        if len(part) >= 80:
            result.append(
                {
                    "id": hashlib.sha1((source["id"] + "#" + str(number) + "#" + part).encode("utf-8")).hexdigest(),
                    "source_id": source["id"],
                    "url": "https://mtpc.edu.vn/tuyen-sinh",
                    "title": source["title"],
                    "text": part,
                    "source_type": source["type"],
                    "source_year": source["year"],
                    "updated_at": source["updated_at"],
                    "origin": source.get("origin", "seed"),
                }
            )
            number += 1
        if end >= len(text):
            break
        offset = max(offset + 1, end - overlap)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source_dir", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--min-year", type=int, default=2024)
    parser.add_argument("--transcriptions", type=Path)
    args = parser.parse_args()
    sources = []
    skipped = []
    errors = []
    for path in sorted(args.source_dir.rglob("*")):
        if not path.is_file() or path.name.startswith(("~$", "~WRL")):
            continue
        if path.suffix.lower() not in SUPPORTED:
            skipped.append(str(path))
            continue
        if any(pattern in path.name.lower() for pattern in PRIVATE_FILE_PATTERNS):
            skipped.append(str(path) + " (private student data)")
            continue
        try:
            text = extract(path)
            if len(text) < 80:
                skipped.append(str(path))
                continue
            year = infer_year(path, text)
            active = path.name.lower() in ACTIVE_DOCUMENTS
            source_id = "manual-" + hashlib.sha1(str(path.relative_to(args.source_dir)).encode("utf-8")).hexdigest()[:16]
            sources.append(
                {
                    "id": source_id,
                    "title": path.stem.replace("_", " ").strip(),
                    "file_name": path.name,
                    "relative_path": str(path.relative_to(args.source_dir)).replace("\\", "/"),
                    "type": infer_type(path.name, text),
                    "year": year,
                    "active": active,
                    "origin": "seed",
                    "warning": "" if active else "Tài liệu lịch sử hoặc ngày hiệu lực chưa được xác nhận; không dùng để trả lời hiện tại.",
                    "updated_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                    "text": text,
                }
            )
        except Exception as error:
            errors.append({"file": str(path), "error": str(error)})
    if args.transcriptions and args.transcriptions.is_file():
        try:
            transcription_data = json.loads(args.transcriptions.read_text(encoding="utf-8"))
            for item in transcription_data.get("sources", []):
                text = clean(str(item.get("text", "")))
                if len(text) < 80 or item.get("private"):
                    continue
                source_id = "image-" + hashlib.sha1(str(item.get("file_name", item.get("title", "image"))).encode("utf-8")).hexdigest()[:16]
                sources.append({
                    "id": source_id,
                    "title": str(item.get("title", "Thông báo tuyển sinh")),
                    "file_name": str(item.get("file_name", "")),
                    "relative_path": str(item.get("file_name", "")),
                    "type": str(item.get("type", "Tuyển sinh")),
                    "year": int(item.get("year", args.min_year)),
                    "active": bool(item.get("active", False)),
                    "origin": "seed",
                    "warning": str(item.get("warning", "")),
                    "updated_at": str(item.get("updated_at", datetime.now(timezone.utc).isoformat())),
                    "sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                    "text": text,
                })
        except Exception as error:
            errors.append({"file": str(args.transcriptions), "error": str(error)})
    chunks = []
    for source in sources:
        if source["active"]:
            chunks.extend(chunks_for(source))
    payload = {
        "version": 2,
        "built_at": datetime.now(timezone.utc).isoformat(),
        "source_count": len(sources),
        "active_source_count": sum(1 for item in sources if item["active"]),
        "chunk_count": len(chunks),
        "sources": sources,
        "chunks": chunks,
        "skipped": skipped,
        "errors": errors,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: payload[key] for key in ("source_count", "active_source_count", "chunk_count", "errors")}, ensure_ascii=False))
    return 0 if not errors else 2


if __name__ == "__main__":
    sys.exit(main())
